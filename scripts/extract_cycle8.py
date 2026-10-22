"""Generate the cycle8 seed and the xlsm fidelity fixture.

Dev-only script. Requires openpyxl:
  /home/UCBerkeley/.venv/bin/python scripts/extract_cycle8.py
"""

from __future__ import annotations

import datetime as dt
import json
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
XLSM = Path('/home/UCBerkeley/Downloads/cycle8.xlsm')
SEED_PATH = ROOT / 'src' / 'data' / 'cycle8.seed.json'
EXPECTED_PATH = ROOT / 'tests' / 'fixtures' / 'cycle8.expected.json'

WEEK_COUNT = 12
WEEK_STRIDE = 37
DAY_DEFS = [
    ('day-mon', 'Пн', 2),
    ('day-wed', 'Ср', 9),
    ('day-fri', 'Пт', 16),
]
GROUP_COLUMNS = [(6, 7, 8, 9), (10, 11, 12, 13), (14, 15, 16, 17), (18, 19, 20, 21)]


def number(value, default: float = 0.0) -> float:
    if value is None or isinstance(value, str):
        return default
    return float(value)


def optional_number(value):
    if value is None or isinstance(value, str):
        return None
    return float(value)


def as_minutes(value) -> float:
    if value is None:
        return 0.0
    if isinstance(value, dt.time):
        return value.hour * 60 + value.minute + value.second / 60
    if isinstance(value, dt.datetime):
        return value.hour * 60 + value.minute + value.second / 60
    if isinstance(value, dt.timedelta):
        return value.total_seconds() / 60
    if isinstance(value, (int, float)):
        # Excel time is a fraction of a day; entry times are already minutes.
        return float(value) * 1440 if 0 <= float(value) < 1 else float(value)
    return 0.0


def rounded(value, digits: int = 6):
    if value is None:
        return None
    return round(float(value), digits)


def main() -> None:
    wb = openpyxl.load_workbook(XLSM, data_only=True)
    ws = wb['Цикл']
    fsheet = wb['F']

    # Exercises are fixed by their first appearance in week 1.
    exercises = []
    exercise_by_name: dict[str, dict] = {}
    for _, _, day_offset in DAY_DEFS:
        start_row = 1 + day_offset
        for row in range(start_row, start_row + 6):
            name = ws.cell(row, 3).value
            if not name:
                continue
            if name in exercise_by_name:
                continue
            exercise = {
                'id': f'ex-{len(exercises) + 1:02d}',
                'name': name,
                'oneRM': optional_number(ws.cell(row, 25).value),
                'impact': number(ws.cell(row, 4).value, 1.0),
                'multiplier': number(ws.cell(row, 5).value, 1.0),
                'rounding': 'auto',
            }
            exercises.append(exercise)
            exercise_by_name[name] = exercise

    if len(exercises) != 18:
        raise RuntimeError(f'Ожидалось 18 упражнений, найдено {len(exercises)}')

    weeks = []
    expected_weeks = []
    entry_count = 0

    for week_index in range(WEEK_COUNT):
        week_start = 1 + week_index * WEEK_STRIDE
        week_id = f'week-{week_index + 1}'
        week_days = []
        expected_days = []

        for day_index, (day_id, day_name, day_offset) in enumerate(DAY_DEFS):
            block_start = week_start + day_offset
            entries = []
            expected_entries = []

            for local_row in range(6):
                row = block_start + local_row
                name = ws.cell(row, 3).value
                if not name:
                    continue
                exercise = exercise_by_name[name]
                groups = []
                weights = []
                for group_index, (weight_col, reps_col, sets_col, percent_col) in enumerate(
                    GROUP_COLUMNS, start=1
                ):
                    percent_raw = ws.cell(row, percent_col).value
                    if percent_raw is None or float(percent_raw) <= 0:
                        continue
                    groups.append(
                        {
                            'id': f'group-w{week_index + 1}-d{day_index + 1}-e{len(entries) + 1}-g{group_index}',
                            'percent': rounded(float(percent_raw) * 100, 2),
                            'reps': int(ws.cell(row, reps_col).value),
                            'sets': int(ws.cell(row, sets_col).value),
                        }
                    )
                    weights.append(optional_number(ws.cell(row, weight_col).value))

                entry_id = f'entry-w{week_index + 1}-d{day_index + 1}-e{len(entries) + 1}'
                entries.append(
                    {
                        'id': entry_id,
                        'exerciseId': exercise['id'],
                        'groups': groups,
                    }
                )
                avg_weight = number(ws.cell(row, 23).value)
                row_one_rm = optional_number(ws.cell(row, 25).value)
                row_multiplier = number(ws.cell(row, 5).value, 1.0)
                relative_intensity = (
                    rounded(avg_weight / (row_one_rm * row_multiplier))
                    if row_one_rm and row_multiplier > 0
                    else None
                )
                expected_entries.append(
                    {
                        'entryId': entry_id,
                        'groups': [{'weight': rounded(w)} for w in weights],
                        'tonnage': rounded(number(ws.cell(row, 22).value)),
                        'kpsh': rounded(number(ws.cell(row, 27).value)),
                        'funtikov': rounded(number(ws.cell(row, 26).value)),
                        'timeMinutes': rounded(as_minutes(ws.cell(row, 28).value)),
                        'avgWeight': rounded(avg_weight),
                        'relativeIntensity': relative_intensity,
                    }
                )
                entry_count += 1

            day_total_row = week_start + day_offset + 6
            week_days.append({'dayTemplateId': day_id, 'entries': entries})
            expected_days.append(
                {
                    'dayTemplateId': day_id,
                    'tonnage': rounded(number(ws.cell(day_total_row, 22).value)),
                    'kpsh': rounded(number(ws.cell(day_total_row, 27).value)),
                    'funtikov': rounded(number(ws.cell(day_total_row, 26).value)),
                    'timeMinutes': rounded(as_minutes(ws.cell(day_total_row, 28).value)),
                    'entries': expected_entries,
                }
            )

        weeks.append(
            {
                'id': week_id,
                'label': f'Неделя {week_index + 1}',
                'isDeload': False,
                'days': week_days,
            }
        )
        expected_weeks.append(
            {
                'weekId': week_id,
                'tonnage': rounded(number(ws.cell(week_start, 22).value)),
                'kpsh': rounded(number(ws.cell(week_start, 27).value)),
                'funtikov': rounded(number(ws.cell(week_start, 26).value)),
                'timeMinutes': rounded(as_minutes(ws.cell(week_start, 28).value)),
                'days': expected_days,
            }
        )

    funtikov_table = [0.0]
    for row in range(1, 101):
        funtikov_table.append(rounded(number(fsheet.cell(row, 2).value)))

    project = {
        'schemaVersion': 1,
        'id': 'cycle8-template',
        'name': 'Цикл 8 — шаблон',
        'createdAt': '2026-10-22T00:00:00.000Z',
        'updatedAt': '2026-10-22T00:00:00.000Z',
        'unit': 'kg',
        'defaultRounding': 'auto',
        'exercises': exercises,
        'dayTemplates': [
            {'id': 'day-mon', 'name': 'Пн'},
            {'id': 'day-wed', 'name': 'Ср'},
            {'id': 'day-fri', 'name': 'Пт'},
        ],
        'weeks': weeks,
        'progression': {
            'mode': 'compound',
            'stepPercent': 1,
            'deload': {'mode': 'hold', 'reductionPercent': 10},
        },
        'metrics': {
            'funtikovTable': funtikov_table,
            'timeThresholds': [
                {'belowPercent': 45, 'factor': 0.5},
                {'belowPercent': 55, 'factor': 0.8},
                {'belowPercent': 65, 'factor': 0.9},
                {'belowPercent': 75, 'factor': 1.0},
                {'belowPercent': 85, 'factor': 1.1},
            ],
            'timeFallbackFactor': 1.2,
            'impactHeavyThreshold': 0.8,
            'impactLightThreshold': 0.4,
            'impactHeavyFactor': 6.0,
            'impactMediumFactor': 3.0,
            'impactLightFactor': 1.5,
        },
    }

    SEED_PATH.parent.mkdir(parents=True, exist_ok=True)
    EXPECTED_PATH.parent.mkdir(parents=True, exist_ok=True)
    SEED_PATH.write_text(json.dumps(project, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    EXPECTED_PATH.write_text(
        json.dumps({'weeks': expected_weeks}, ensure_ascii=False, indent=2) + '\n',
        encoding='utf-8',
    )

    print(f'weeks={len(weeks)} exercises={len(exercises)} entries={entry_count}')
    print(f'wrote {SEED_PATH}')
    print(f'wrote {EXPECTED_PATH}')
    print(
        'week1:',
        expected_weeks[0]['tonnage'],
        expected_weeks[0]['kpsh'],
        expected_weeks[0]['funtikov'],
        expected_weeks[0]['timeMinutes'],
    )


if __name__ == '__main__':
    main()
